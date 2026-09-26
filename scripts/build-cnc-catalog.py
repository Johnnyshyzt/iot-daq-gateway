#!/usr/bin/env python3
"""Build the versioned CNC catalog from docs/catalog/cnc-data-template.xlsx.

Sheet1 is one column per brand. Sheet2 is controller series. Synonyms collapse
into canonical items. A name that is not a synonym becomes a brand-specific item
with a stable English id. Run-state and work-mode stay separate: several brands
list both on the same machine.
"""

from __future__ import annotations

import json
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "docs" / "catalog" / "cnc-data-template.xlsx"
OUT = ROOT / "docs" / "catalog" / "cnc-catalog.json"
SEED = ROOT / "data" / "seed" / "point-templates"
VERSION = "2026.09.26"

# (id, nameZh, dataType, unit, category, synonyms)
CANONICAL = [
    ("state", "运行状态", "string", "", "state", ["运行状态", "状态"]),
    ("workMode", "工作模式", "string", "", "state", ["工作模式", "操作模式", "工作方式", "模式", "执行模式"]),
    ("alarm", "报警信息", "string", "", "alarm", ["报警信息", "告警状态", "是否报警"]),
    ("alarmNumber", "警报号", "string", "", "alarm", ["警报号"]),
    ("estop", "急停状态", "bool", "", "state", ["急停状态", "是否急停"]),
    ("isRunning", "是否运行", "bool", "", "state", ["是否运行"]),
    ("isReset", "是否复位", "bool", "", "state", ["是否复位"]),
    ("isCutting", "是否切削", "bool", "", "state", ["是否切削"]),
    ("program", "当前程序", "string", "", "program", [
        "当前程序号", "当前执行程序名", "当前程序名称", "当前程序名", "当前NC程序名", "运行程序编号",
        "加工程序", "程序名称", "CNC 程序名", "程序名", "程序号",
    ]),
    ("programMain", "主程序", "string", "", "program", [
        "主程序号", "主程序名", "主程序名称", "当前主程序号", "选择程序编号",
    ]),
    ("programComment", "程序注释", "string", "", "program", [
        "主程序备注", "当前程序备注", "程序注释", "程序名备注", "主程序注释", "当前程序注释",
    ]),
    ("programBlock", "当前程序段", "string", "", "program", [
        "当前执行代码", "当前程序段内容", "执行代码段", "加工代码", "G代码",
    ]),
    ("programLine", "程序行号", "number", "", "program", [
        "程序行号", "执行行号", "当前程序行号", "当前NC程序段号", "当前程序块号",
    ]),
    ("programPath", "程序路径", "string", "", "program", ["执行程序路径"]),
    ("partCount", "加工件数", "number", "pcs", "counter", [
        "加工个数", "加工件数", "工件数", "产量", "加工数量", "当前加工件数", "单批件数",
        "完成加工数", "零件计数1",
    ]),
    ("partCountTotal", "总加工件数", "number", "pcs", "counter", [
        "总加工个数", "总加工件数", "总工件数", "总件数", "零件计数2",
    ]),
    ("partCountTarget", "目标件数", "number", "pcs", "counter", [
        "需求工件数", "单批目标件数", "总目标件数", "目标加工数",
    ]),
    ("cycleTime", "循环时间", "number", "s", "time", ["循环时间", "最后一次循环时间", "循环时间毫秒值"]),
    ("powerOnTime", "上电时间", "number", "s", "time", [
        "上电时间", "开机时间", "累计开机时间", "机床上电时间", "总通电时间", "系统启动时间",
    ]),
    ("runTime", "运行时间", "number", "s", "time", [
        "运行时间", "当前运行时间", "机床运行时间", "程序运行时间", "总移动时间",
    ]),
    ("cutTime", "切削时间", "number", "s", "time", ["切削时间", "加工时间", "当前切削时间", "累计加工时间"]),
    ("remainTime", "剩余加工时间", "number", "s", "time", ["剩余加工时间"]),
    ("spindleSpeed", "主轴转速", "number", "rpm", "spindle", [
        "主轴转速", "实际主轴转速", "主轴实际转速", "实际主轴速度",
    ]),
    ("spindleSpeedCmd", "设定主轴转速", "number", "rpm", "spindle", [
        "设定主轴转速", "主轴设定转速", "编程主轴转速", "主轴编程转速", "设定主轴速度",
    ]),
    ("spindleOverride", "主轴倍率", "number", "%", "spindle", ["主轴倍率"]),
    ("spindleLoad", "主轴负载", "number", "%", "spindle", ["主轴负载"]),
    ("spindleTemp", "主轴温度", "number", "°C", "temperature", ["主轴温度"]),
    ("feedRate", "进给速度", "number", "mm/min", "feed", ["进给速度", "实际进给速度", "进给实际速度"]),
    ("feedRateCmd", "设定进给速度", "number", "mm/min", "feed", [
        "设定进给速度", "编程进给速度", "G94设定进给速率", "进给设定速度",
    ]),
    ("feedOverride", "进给倍率", "number", "%", "feed", ["进给倍率", "当前进给倍率", "前进的覆盖"]),
    ("rapidOverride", "快速倍率", "number", "%", "feed", [
        "快速倍率", "快进倍率", "当前快速倍率", "快速进给倍率", "快速移动倍率",
    ]),
    ("machinePosition", "机械坐标", "string", "mm", "axis", [
        "机械坐标", "机械位置", "机床坐标", "轴位置", "坐标", "当前坐标", "坐标信息", "实际坐标",
    ]),
    ("relativePosition", "相对坐标", "string", "mm", "axis", ["相对坐标", "相对位置"]),
    ("absolutePosition", "绝对坐标", "string", "mm", "axis", ["绝对坐标", "绝对位置", "当前位置"]),
    ("distanceToGo", "剩余坐标", "string", "mm", "axis", ["剩余坐标", "剩余位置"]),
    ("machinePositionX", "机械坐标X", "number", "mm", "axis", ["机械坐标X"]),
    ("machinePositionY", "机械坐标Y", "number", "mm", "axis", ["机械坐标Y"]),
    ("machinePositionZ", "机械坐标Z", "number", "mm", "axis", ["机械坐标Z"]),
    ("absolutePositionX", "绝对坐标X", "number", "mm", "axis", ["绝对坐标X"]),
    ("absolutePositionY", "绝对坐标Y", "number", "mm", "axis", ["绝对坐标Y"]),
    ("absolutePositionZ", "绝对坐标Z", "number", "mm", "axis", ["绝对坐标Z"]),
    ("relativePositionX", "相对坐标X", "number", "mm", "axis", ["相对坐标X"]),
    ("relativePositionY", "相对坐标Y", "number", "mm", "axis", ["相对坐标Y"]),
    ("relativePositionZ", "相对坐标Z", "number", "mm", "axis", ["相对坐标Z"]),
    ("distanceToGoX", "剩余距离X", "number", "mm", "axis", ["剩余距离X"]),
    ("distanceToGoY", "剩余距离Y", "number", "mm", "axis", ["剩余距离Y"]),
    ("distanceToGoZ", "剩余距离Z", "number", "mm", "axis", ["剩余距离Z"]),
    ("axisLoadX", "X轴负载", "number", "%", "load", ["x轴负载", "X轴负载"]),
    ("axisLoadY", "Y轴负载", "number", "%", "load", ["y轴负载", "Y轴负载"]),
    ("axisLoadZ", "Z轴负载", "number", "%", "load", ["z轴负载", "Z轴负载"]),
    ("toolNumber", "刀具号", "string", "", "tool", ["刀具号", "当前刀具", "刀具编号", "当前刀具刀补号", "主轴刀号"]),
    ("toolNext", "下一把刀", "string", "", "tool", ["预选刀具", "下一把刀具号"]),
    ("serialNumber", "序列号", "string", "", "system-info", ["序列号", "机床序列号", "系统ID", "唯一标识"]),
    ("softwareVersion", "软件版本", "string", "", "system-info", ["软件版本", "版本号", "版本信息"]),
    ("systemType", "系统类型", "string", "", "system-info", ["系统类型", "设备类型"]),
    ("axisCount", "有效轴数", "number", "", "system-info", ["有效轴数", "伺服轴数", "CNC 当前有效轴数"]),
    ("servoLoad", "伺服负载", "number", "%", "load", ["伺服负载", "负载"]),
]

# Names that are not synonyms of a canonical item. id is stable English.
SPECIFIC = {
    "风扇转速": ("fanSpeed", "number", "rpm", "temperature"),
    "x轴温度": ("axisTempX", "number", "°C", "temperature"),
    "y轴温度": ("axisTempY", "number", "°C", "temperature"),
    "z轴温度": ("axisTempZ", "number", "°C", "temperature"),
    "操作信息": ("operatorMessage", "string", "", "system-info"),
    "剩余加工时间": ("remainMachiningTime", "number", "s", "time"),  # already canonical remainTime — unused if synonym hits first
    "电压": ("voltage", "number", "V", "load"),
    "电流": ("current", "number", "A", "load"),
    "电机温度": ("motorTemp", "number", "°C", "temperature"),
    "驱动温度": ("driveTemp", "number", "°C", "temperature"),
    "设定进给超限": ("feedLimitCmd", "number", "%", "feed"),
    "设定主轴超限": ("spindleLimitCmd", "number", "%", "spindle"),
    "工作笔数": ("workStrokeCount", "number", "pcs", "counter"),
    "伺服轴名称": ("servoAxisNames", "string", "", "axis"),
    "门阵版本": ("fpgaVersion", "string", "", "system-info"),
    "梯图版本": ("ladderVersion", "string", "", "system-info"),
    "制造商": ("manufacturer", "string", "", "system-info"),
    "主轴上电时间": ("spindlePowerOnTime", "number", "s", "time"),
    "主轴运行时间": ("spindleRunTime", "number", "s", "time"),
    "报价信息": ("quoteInfo", "string", "", "system-info"),
    "DNC模式": ("dncMode", "string", "", "state"),
    "刀具信息": ("toolInfo", "string", "", "tool"),
    "编辑进给": ("editedFeed", "number", "mm/min", "feed"),
    "最大进给速度": ("maxFeedRate", "number", "mm/min", "feed"),
    "轮廓进给速度": ("contourFeedRate", "number", "mm/min", "feed"),
    "驱动单元电流": ("driveCurrent", "number", "A", "load"),
    "负载电流": ("loadCurrent", "number", "A", "load"),
    "额定电流": ("ratedCurrent", "number", "A", "load"),
    "X轴实际进给速率": ("feedRateActualX", "number", "mm/min", "feed"),
    "Y轴实际进给速率": ("feedRateActualY", "number", "mm/min", "feed"),
    "Z轴实际进给速率": ("feedRateActualZ", "number", "mm/min", "feed"),
    "X轴进给率": ("feedOverrideX", "number", "%", "feed"),
    "Y轴进给率": ("feedOverrideY", "number", "%", "feed"),
    "Z轴进给率": ("feedOverrideZ", "number", "%", "feed"),
    "X轴实际位置": ("actualPositionX", "number", "mm", "axis"),
    "Y轴实际位置": ("actualPositionY", "number", "mm", "axis"),
    "Z轴实际位置": ("actualPositionZ", "number", "mm", "axis"),
    "进给速度超程倍率": ("feedOvertravelOverride", "number", "%", "feed"),
    "刀槽号": ("toolPot", "string", "", "tool"),
    "机床型号": ("machineModel", "string", "", "system-info"),
    "换刀次数": ("toolChangeCount", "number", "pcs", "counter"),
    "上一次循环时间": ("previousCycleTime", "number", "s", "time"),
    "手轮倍率": ("handwheelOverride", "number", "%", "feed"),
    "当前手动倍率": ("manualOverride", "number", "%", "feed"),
    "当前手轮倍率": ("mpgOverride", "number", "%", "feed"),
    "循环时间小时值": ("cycleTimeHours", "number", "h", "time"),
    "加工时间毫秒值": ("machiningTimeMs", "number", "ms", "time"),
    "加工时间小时值": ("machiningTimeHours", "number", "h", "time"),
    "警告号（": ("warningNumber", "string", "", "alarm"),
    "M代码": ("mCode", "string", "", "program"),
    "G 模态信息": ("gModal", "string", "", "program"),
    "M 模态信息": ("mModal", "string", "", "program"),
    "G00速度": ("g00Speed", "number", "mm/min", "feed"),
    "G00运行时间": ("g00RunTime", "number", "s", "time"),
    "G01运行时间": ("g01RunTime", "number", "s", "time"),
    "程序坐标": ("programPosition", "string", "mm", "axis"),
    "程序位置": ("programPosition", "string", "mm", "axis"),
    "轴名称": ("axisNames", "string", "", "axis"),
    "指令速度": ("commandVelocity", "number", "mm/min", "feed"),
    "实际速度": ("actualVelocity", "number", "mm/min", "feed"),
    "实际加速度": ("actualAcceleration", "number", "mm/s2", "axis"),
    "实际加加速度": ("actualJerk", "number", "mm/s3", "axis"),
    "反向跃冲补偿值": ("backlashJumpComp", "number", "mm", "axis"),
    "逻辑轴补偿后机床坐标系坐标": ("logicalAxisCompensated", "string", "mm", "axis"),
    "逻辑轴补偿前机床坐标系坐标": ("logicalAxisRaw", "string", "mm", "axis"),
    "逻辑轴双驱轴误差1": ("gantryError1", "number", "mm", "axis"),
    "逻辑轴双驱轴误差2": ("gantryError2", "number", "mm", "axis"),
    "逻辑轴双驱轴误差3": ("gantryError3", "number", "mm", "axis"),
    "轮廓误差": ("contourError", "number", "mm", "axis"),
    "刀尖点指令速度": ("toolTipCommandSpeed", "number", "mm/min", "tool"),
    "刀头点指令速度": ("toolHeadCommandSpeed", "number", "mm/min", "tool"),
    "驱动运行状态": ("driveState", "string", "", "state"),
    "实际反馈速度": ("feedbackSpeed", "number", "mm/min", "axis"),
    "实际励磁电流": ("fieldCurrent", "number", "A", "load"),
    "实际转矩(推力)电流": ("torqueCurrent", "number", "A", "load"),
    "实际输出转矩": ("outputTorque", "number", "Nm", "load"),
    "电流负载率": ("currentLoadRate", "number", "%", "load"),
    "电流负载率峰值": ("currentLoadPeak", "number", "%", "load"),
    "速度跟踪误差": ("speedTrackError", "number", "mm/min", "axis"),
    "转矩(推力)负载率": ("torqueLoadRate", "number", "%", "load"),
    "刀具列表": ("toolList", "string", "", "tool"),
    "自动启动时间": ("autoStartTime", "number", "s", "time"),
    "暂停时间": ("dwellTime", "number", "s", "time"),
    "刀库刀号": ("magazineTool", "string", "", "tool"),
    "偏移坐标": ("offsetPosition", "string", "mm", "axis"),
    "主轴电流": ("spindleCurrent", "number", "A", "spindle"),
    "主轴功率": ("spindlePower", "number", "kW", "spindle"),
    "主轴转矩": ("spindleTorque", "number", "Nm", "spindle"),
    "主轴1电流负载率": ("spindle1CurrentLoad", "number", "%", "spindle"),
    "主轴1功率负载率": ("spindle1PowerLoad", "number", "%", "spindle"),
    "主轴1转矩负载率": ("spindle1TorqueLoad", "number", "%", "spindle"),
    "电控柜温度": ("cabinetTemp", "number", "°C", "temperature"),
    "系统信息": ("systemInfo", "string", "", "system-info"),
}

BRANDS = [
    ("fanuc", "发那科", "Fanuc", "发那科"),
    ("siemens", "西门子", "Siemens", "西门子"),
    ("mitsubishi", "三菱", "Mitsubishi", "三菱"),
    ("syntec", "新代", "Syntec", "新代"),
    ("gsk", "广数", "GSK", "广数"),
    ("mazak-smart", "马扎克 smart/smooth", "Mazak smart/smooth", "马扎克（smart/smooth）"),
    ("mazak-matrix", "马扎克 640/matrix", "Mazak 640/matrix", "马扎克（640/matrix）"),
    ("knd", "凯恩帝", "KND", "凯恩第"),
    ("heidenhain", "海德汉", "Heidenhain", "海德汉"),
    ("hnc", "华中数控", "Huazhong HNC", "华中数控"),
    ("fagor", "法格", "Fagor", "法格"),
    ("brother", "兄弟", "Brother", "兄弟"),
    ("haas", "哈斯", "Haas", "哈斯"),
    ("baoyuan", "宝元", "Baoyuan/LNC", "宝元"),
    ("kede", "科德", "KEDE", "科德"),
    ("citizen", "三菱西铁城", "Citizen (Mitsubishi)", "三菱西铁城"),
    ("dmg-mori", "三菱森精机", "DMG Mori (Mitsubishi)", "三菱森精机"),
    ("jdsoft", "精雕", "JDSoft", "精雕"),
    ("delta", "台达", "Delta", "台达"),
]

# Sheet2 column -> brand id. Mazak is split by model name.
SHEET2_BRAND = {
    "三菱": "mitsubishi",
    "发那科": "fanuc",
    "西门子": "siemens",
    "新代": "syntec",
    "广数": "gsk",
    "海德汉": "heidenhain",
    "华中": "hnc",
    "凯恩帝": "knd",
    "科德": "kede",
    "精雕": "jdsoft",
    "宝元": "baoyuan",
}

MAZAK_MODEL_BRAND = {
    "smart": "mazak-smart",
    "smooth": "mazak-smart",
    "640": "mazak-matrix",
    "matrix": "mazak-matrix",
}

MODEL_IDS = {
    ("fanuc", "全系列网口"): "fanuc-ethernet",
    ("fanuc", "无网口且支持扩展PCMCIA网口系列"): "fanuc-pcmcia",
    ("siemens", "828D"): "siemens-828d",
    ("siemens", "840dsl"): "siemens-840dsl",
    ("siemens", "808D"): "siemens-808d",
    ("siemens", "840d(PCU5.0)"): "siemens-840d-pcu50",
    ("siemens", "810d(PCU5.0)"): "siemens-810d-pcu50",
    ("mitsubishi", "M70"): "mitsubishi-m70",
    ("mitsubishi", "M80"): "mitsubishi-m80",
    ("mitsubishi", "E70"): "mitsubishi-e70",
    ("mitsubishi", "E80"): "mitsubishi-e80",
    ("syntec", "116.1X版本以上"): "syntec-116",
    ("gsk", "988ta"): "gsk-988ta",
    ("mazak-smart", "smart"): "mazak-smart",
    ("mazak-smart", "smooth"): "mazak-smooth",
    ("mazak-matrix", "640"): "mazak-640",
    ("mazak-matrix", "matrix"): "mazak-matrix",
    ("knd", "5.0版本以上"): "knd-5",
    ("heidenhain", "530"): "heidenhain-530",
    ("heidenhain", "640"): "heidenhain-640",
    ("heidenhain", "620"): "heidenhain-620",
    ("hnc", "8系"): "hnc-8",
    ("kede", "GNC60"): "kede-gnc60",
    ("kede", "GNC62"): "kede-gnc62",
    ("jdsoft", "JD50"): "jdsoft-jd50",
    ("jdsoft", "JD60"): "jdsoft-jd60",
    ("baoyuan", "T5800系列"): "baoyuan-t5800",
}

PROTOCOLS = {
    "fanuc": [("fanuc.fake", "simulator", "fake", 1, "发那科模拟（兼容三态）"),
              ("fanuc.sim", "simulator", "sim", 1, "发那科模拟器"),
              ("fanuc.focas", "driver", "focas", 1, "发那科 FOCAS")],
    "siemens": [("siemens.sim", "simulator", "sim", 1, "西门子模拟器"),
                ("siemens.opcua", "stub", "opcua", 2, "西门子 OPC UA")],
    "mitsubishi": [("mitsubishi.sim", "simulator", "sim", 1, "三菱模拟器"),
                   ("mitsubishi.custom", "stub", "mitsubishi", 2, "三菱协议")],
    "syntec": [("syntec.sim", "simulator", "sim", 1, "新代模拟器"),
               ("syntec.custom", "stub", "syntec", 2, "新代协议")],
    "gsk": [("gsk.sim", "simulator", "sim", 1, "广数模拟器"),
            ("gsk.custom", "stub", "gsk", 2, "广数协议")],
    "mazak-smart": [("mazak-smart.sim", "simulator", "sim", 1, "马扎克 smart/smooth 模拟器"),
                    ("mazak-smart.mtconnect", "stub", "mtconnect", 2, "马扎克 MTConnect")],
    "mazak-matrix": [("mazak-matrix.sim", "simulator", "sim", 1, "马扎克 640/matrix 模拟器"),
                     ("mazak-matrix.mtconnect", "stub", "mtconnect", 2, "马扎克 MTConnect")],
    "knd": [("knd.sim", "simulator", "sim", 1, "凯恩帝模拟器"),
            ("knd.custom", "stub", "knd", 2, "凯恩帝协议")],
    "heidenhain": [("heidenhain.sim", "simulator", "sim", 1, "海德汉模拟器"),
                   ("heidenhain.lsv2", "stub", "lsv2", 2, "海德汉 LSV2")],
    "hnc": [("hnc.sim", "simulator", "sim", 1, "华中数控模拟器"),
            ("hnc.custom", "stub", "hnc", 2, "华中数控协议")],
    "fagor": [("fagor.sim", "simulator", "sim", 1, "法格模拟器"),
              ("fagor.custom", "stub", "fagor", 2, "法格协议")],
    "brother": [("brother.sim", "simulator", "sim", 1, "兄弟模拟器"),
                ("brother.custom", "stub", "brother", 2, "兄弟协议")],
    "haas": [("haas.sim", "simulator", "sim", 1, "哈斯模拟器"),
             ("haas.mtconnect", "stub", "mtconnect", 2, "哈斯 MTConnect")],
    "baoyuan": [("baoyuan.sim", "simulator", "sim", 1, "宝元模拟器"),
                ("baoyuan.custom", "stub", "baoyuan", 2, "宝元协议")],
    "kede": [("kede.sim", "simulator", "sim", 1, "科德模拟器"),
             ("kede.custom", "stub", "kede", 2, "科德协议")],
    "citizen": [("citizen.sim", "simulator", "sim", 1, "西铁城模拟器"),
                ("citizen.mitsubishi", "stub", "mitsubishi", 2, "西铁城（三菱）协议")],
    "dmg-mori": [("dmg-mori.sim", "simulator", "sim", 1, "森精机模拟器"),
                 ("dmg-mori.mitsubishi", "stub", "mitsubishi", 2, "森精机（三菱）协议")],
    "jdsoft": [("jdsoft.sim", "simulator", "sim", 1, "精雕模拟器"),
               ("jdsoft.custom", "stub", "jdsoft", 2, "精雕协议")],
    "delta": [("delta.sim", "simulator", "sim", 1, "台达模拟器"),
              ("delta.custom", "stub", "delta", 2, "台达协议")],
}

SIM_PARAMS = [
    {"name": "host", "type": "string", "label": "主机", "required": True, "default": "127.0.0.1"},
    {"name": "port", "type": "int", "label": "端口", "required": True, "default": 8193},
    {"name": "timeoutMs", "type": "int", "label": "超时 (ms)", "required": False, "default": 3000},
]

PROTOCOL_PARAMS = {
    "fake": SIM_PARAMS,
    "sim": SIM_PARAMS,
    "focas": [
        {"name": "host", "type": "string", "label": "主机", "required": True, "default": "192.168.1.10"},
        {"name": "port", "type": "int", "label": "端口", "required": True, "default": 8193},
        {"name": "timeoutMs", "type": "int", "label": "FOCAS 超时 (ms)", "required": True, "default": 3000},
    ],
    "opcua": [
        {"name": "host", "type": "string", "label": "主机", "required": True, "default": "192.168.1.20"},
        {"name": "port", "type": "int", "label": "端口", "required": True, "default": 4840},
        {"name": "timeoutMs", "type": "int", "label": "超时 (ms)", "required": False, "default": 5000},
        {"name": "path", "type": "string", "label": "端点路径", "required": False, "default": ""},
        {"name": "namespace", "type": "string", "label": "命名空间", "required": False, "default": ""},
    ],
    "mtconnect": [
        {"name": "host", "type": "string", "label": "主机", "required": True, "default": "192.168.1.30"},
        {"name": "port", "type": "int", "label": "端口", "required": True, "default": 5000},
        {"name": "timeoutMs", "type": "int", "label": "超时 (ms)", "required": False, "default": 5000},
        {"name": "path", "type": "string", "label": "路径", "required": False, "default": "/current"},
    ],
    "lsv2": [
        {"name": "host", "type": "string", "label": "主机", "required": True, "default": "192.168.1.40"},
        {"name": "port", "type": "int", "label": "端口", "required": True, "default": 19000},
        {"name": "timeoutMs", "type": "int", "label": "超时 (ms)", "required": False, "default": 5000},
    ],
    "ftp": [
        {"name": "host", "type": "string", "label": "主机", "required": True, "default": "192.168.1.50"},
        {"name": "port", "type": "int", "label": "端口", "required": True, "default": 21},
        {"name": "timeoutMs", "type": "int", "label": "超时 (ms)", "required": False, "default": 8000},
        {"name": "path", "type": "string", "label": "路径或共享", "required": False, "default": "/"},
    ],
}

DEFAULT_CUSTOM = [
    {"name": "host", "type": "string", "label": "主机", "required": True, "default": "192.168.1.60"},
    {"name": "port", "type": "int", "label": "端口", "required": True, "default": 8193},
    {"name": "timeoutMs", "type": "int", "label": "超时 (ms)", "required": False, "default": 3000},
    {"name": "path", "type": "string", "label": "路径", "required": False, "default": ""},
]


def synonym_index():
    index = {}
    items = []
    for item_id, name_zh, data_type, unit, category, synonyms in CANONICAL:
        items.append({
            "id": item_id,
            "nameZh": name_zh,
            "dataType": data_type,
            "unit": unit,
            "category": category,
            "brandSpecific": False,
            "brandId": None,
        })
        for name in synonyms:
            if name in index:
                raise SystemExit(f"synonym {name} mapped twice")
            index[name] = item_id
    return items, index


def yaml_quote(value: str) -> str:
    if value == "" or any(ch in value for ch in ":#{}[]&*!|>'\"%@`"):
        return '"' + value.replace("\\", "\\\\").replace('"', '\\"') + '"'
    return value


def write_template(brand_id: str, display: str, template_id: str, points: list[dict]) -> None:
    lines = [
        "apiVersion: daq.gateway/v1",
        "kind: PointTemplate",
        "metadata:",
        f"  id: {template_id}",
        f"  displayName: {yaml_quote(display)}",
        "spec:",
        f"  adapter: {brand_id}",
        "  points:",
    ]
    for point in points:
        lines.append(f"    - id: {point['id']}")
        lines.append(f"      address: {point['address']}")
        lines.append(f"      dataType: {point['dataType']}")
        unit = point.get("unit") or ""
        lines.append(f"      unit: {yaml_quote(unit)}")
        lines.append("      scale: 1")
        lines.append("      deadband: 0")
        lines.append("      enabled: true")
    SEED.mkdir(parents=True, exist_ok=True)
    (SEED / f"{template_id}.yaml").write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    items, index = synonym_index()
    item_by_id = {item["id"]: item for item in items}
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    sheet1 = wb["Sheet1"]
    headers = {}
    for col in range(1, sheet1.max_column + 1):
        header = sheet1.cell(1, col).value
        if header:
            headers[str(header).strip()] = col

    brands = []
    missing = []
    for brand_id, name_zh, name_en, sheet_header in BRANDS:
        col = headers.get(sheet_header)
        if col is None:
            raise SystemExit(f"missing sheet1 column {sheet_header}")
        mappings = []
        seen_ids = []
        for row in range(2, sheet1.max_row + 1):
            raw = sheet1.cell(row, col).value
            if raw is None:
                continue
            source = str(raw).strip()
            if not source:
                continue
            if source in index:
                item_id = index[source]
                specific = False
            elif source in SPECIFIC:
                short, data_type, unit, category = SPECIFIC[source]
                item_id = f"{brand_id}_{short}"
                specific = True
                if item_id not in item_by_id:
                    item = {
                        "id": item_id,
                        "nameZh": source,
                        "dataType": data_type,
                        "unit": unit,
                        "category": category,
                        "brandSpecific": True,
                        "brandId": brand_id,
                    }
                    items.append(item)
                    item_by_id[item_id] = item
            else:
                missing.append((brand_id, source))
                continue
            mappings.append({"source": source, "itemId": item_id, "brandSpecific": specific})
            if item_id not in seen_ids:
                seen_ids.append(item_id)
        brands.append({
            "id": brand_id,
            "nameZh": name_zh,
            "nameEn": name_en,
            "sheetHeader": sheet_header,
            "mappings": mappings,
            "itemIds": seen_ids,
        })

    if missing:
        for brand_id, source in missing:
            print(f"UNMAPPED {brand_id}: {source}")
        raise SystemExit(f"{len(missing)} unmapped names")

    sheet2 = wb["Sheet2"]
    models_by_brand = {brand["id"]: [] for brand in brands}
    for col in range(1, sheet2.max_column + 1):
        header = sheet2.cell(1, col).value
        if not header:
            continue
        header = str(header).strip()
        names = []
        for row in range(2, sheet2.max_row + 1):
            raw = sheet2.cell(row, col).value
            if raw is None:
                continue
            text = str(raw).strip()
            if text:
                names.append(text)
        if header == "马扎克":
            for name in names:
                owner = MAZAK_MODEL_BRAND.get(name.lower()) or MAZAK_MODEL_BRAND.get(name)
                if owner is None:
                    raise SystemExit(f"mazak model not split: {name}")
                models_by_brand[owner].append(name)
            continue
        if header.startswith("Other"):
            continue
        owner = SHEET2_BRAND.get(header)
        if owner is None:
            raise SystemExit(f"sheet2 column not mapped: {header}")
        models_by_brand[owner].extend(names)

    for brand in brands:
        models = []
        for name in models_by_brand[brand["id"]]:
            key = (brand["id"], name)
            if key not in MODEL_IDS:
                raise SystemExit(f"missing model id for {key}")
            models.append({"id": MODEL_IDS[key], "name": name})
        brand["models"] = models
        adapters = []
        for adapter_id, kind, protocol, phase, display in PROTOCOLS[brand["id"]]:
            params = PROTOCOL_PARAMS.get(protocol, DEFAULT_CUSTOM)
            adapters.append({
                "id": adapter_id,
                "brandId": brand["id"],
                "kind": kind,
                "protocol": protocol,
                "phase": phase,
                "displayName": display,
                "parameters": params,
            })
        brand["adapters"] = adapters

    generic = {
        "id": "generic.ftp",
        "brandId": None,
        "kind": "stub",
        "protocol": "ftp",
        "phase": 2,
        "displayName": "通用 FTP / 共享",
        "parameters": PROTOCOL_PARAMS["ftp"],
        "note": "Sheet2「Other其他」列为 ftp/共享，不是第 20 个品牌。任何品牌都可以改用这个适配器，第二阶段再实现读取。",
    }

    catalog = {
        "version": VERSION,
        "source": "docs/catalog/cnc-data-template.xlsx",
        "notes": [
            "运行状态与工作模式/操作模式保持为两个标准项。同一品牌经常同时给出这两列，合并会丢掉一列。",
            "主轴转速与实际主轴转速合并为 spindleSpeed；设定值单独为 spindleSpeedCmd。",
            "机械坐标与机械位置合并为 machinePosition。",
            "报警信息、告警状态、是否报警合并为 alarm。",
            "加工件数、工件数、产量合并为 partCount。",
            "fanuc-standard 仍是历史三态模板（state/alarm/program）。fanuc-catalog 才是发那科全量目录。",
            "没有出现在 Sheet2 的品牌（兄弟、哈斯、法格、台达、西铁城、森精机）型号列表为空，设备上的控制器型号可留空。",
        ],
        "items": items,
        "brands": brands,
        "genericAdapters": [generic],
    }
    OUT.write_text(json.dumps(catalog, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    for brand in brands:
        points = []
        for item_id in brand["itemIds"]:
            item = item_by_id[item_id]
            if brand["id"] == "fanuc" and item_id == "state":
                address = "cnc/statinfo"
            elif brand["id"] == "fanuc" and item_id == "alarm":
                address = "cnc/alarm"
            elif brand["id"] == "fanuc" and item_id == "program":
                address = "cnc/program"
            else:
                address = f"catalog/{item_id}"
            points.append({
                "id": item_id,
                "address": address,
                "dataType": item["dataType"],
                "unit": item["unit"],
            })
        template_id = "fanuc-catalog" if brand["id"] == "fanuc" else f"{brand['id']}-standard"
        display = f"{brand['nameZh']} 标准目录"
        write_template(brand["id"], display, template_id, points)

    print(f"items={len(items)} brands={len(brands)} -> {OUT}")


if __name__ == "__main__":
    main()
